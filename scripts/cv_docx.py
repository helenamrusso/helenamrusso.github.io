"""Turn the generated cv.html into a Word document (.docx).

Used by build.py; needs `pip install python-docx`. It understands the markup
written in templates/cv.html and the generated publication/presentation lists,
so the .docx always matches the PDF.
"""
from html.parser import HTMLParser

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE

ACCENT = RGBColor(0x1F, 0x4F, 0x8F)
INK = RGBColor(0x1B, 0x22, 0x30)
MUTED = RGBColor(0x5B, 0x64, 0x75)
WARM = RGBColor(0xB4, 0x68, 0x2C)
VOID = {"br", "img", "meta", "link", "hr", "input"}


# ---------------------------------------------------------------- tiny DOM


class Node:
    def __init__(self, tag, attrs=None, parent=None):
        self.tag, self.attrs, self.parent, self.children = tag, dict(attrs or {}), parent, []

    @property
    def classes(self):
        return self.attrs.get("class", "").split()

    def kids(self, tag=None):
        return [c for c in self.children if isinstance(c, Node) and (tag is None or c.tag == tag)]

    def find(self, tag=None, cls=None):
        for c in self.kids():
            if (tag is None or c.tag == tag) and (cls is None or cls in c.classes):
                return c
            hit = c.find(tag, cls)
            if hit:
                return hit
        return None

    def text(self):
        return "".join(c if isinstance(c, str) else c.text() for c in self.children)


class TreeBuilder(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = self.cur = Node("root")

    def handle_starttag(self, tag, attrs):
        node = Node(tag, attrs, self.cur)
        self.cur.children.append(node)
        if tag not in VOID:
            self.cur = node

    def handle_endtag(self, tag):
        n = self.cur
        while n is not self.root and n.tag != tag:
            n = n.parent
        if n is not self.root:
            self.cur = n.parent

    def handle_data(self, data):
        self.cur.children.append(data)


# ---------------------------------------------------------------- docx helpers


def bottom_border(paragraph, color="1F4F8F", size=8):
    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", str(size)), ("w:space", "2"), ("w:color", color)):
        b.set(qn(k), v)
    bdr.append(b)
    ppr.append(bdr)


def add_hyperlink(paragraph, url, text, fmt):
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F4F8F")
    rpr.append(color)
    for flag in ("b", "i"):
        if fmt.get(flag):
            rpr.append(OxmlElement(f"w:{flag}"))
    if fmt.get("size"):
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(fmt["size"] * 2)))
        rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def squash(s):
    return " ".join(s.split())


def inline(paragraph, node, fmt):
    """Write the inline content of node into paragraph, keeping bold/italic/superscript/links."""
    for child in node.children:
        if isinstance(child, str):
            text = squash(child)
            if not text:
                last = getattr(paragraph, "last_char", "")
                if child and last and last not in " \t\n":  # whitespace between two inline elements
                    paragraph.add_run(" ")
                    paragraph.last_char = " "
                continue
            last = getattr(paragraph, "last_char", "")
            if child[:1].isspace() and last and last not in " \t\n":
                text = " " + text
            if child[-1:].isspace():
                text += " "
            paragraph.last_char = text[-1]
            run = paragraph.add_run(text)
            run.bold = fmt.get("b")
            run.italic = fmt.get("i")
            run.font.superscript = fmt.get("sup")
            if fmt.get("size"):
                run.font.size = Pt(fmt["size"])
            run.font.color.rgb = fmt.get("color", INK)
            continue
        tag = child.tag
        if tag == "br":
            paragraph.add_run().add_break()
            paragraph.last_char = "\n"
        elif tag in ("b", "strong"):
            inline(paragraph, child, {**fmt, "b": True, "color": INK})
        elif tag in ("i", "em"):
            inline(paragraph, child, {**fmt, "i": True})
        elif tag == "sup":
            inline(paragraph, child, {**fmt, "sup": True, "color": WARM, "b": True})
        elif tag == "a" and child.attrs.get("href", "").startswith(("http", "mailto")):
            add_hyperlink(paragraph, child.attrs["href"], squash(child.text()), fmt)
            paragraph.last_char = squash(child.text())[-1:] or getattr(paragraph, "last_char", "")
        elif "t" in child.classes:  # publication title
            inline(paragraph, child, {**fmt, "b": True, "color": INK})
        else:
            inline(paragraph, child, fmt)


class Writer:
    def __init__(self):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Inches(0.75))
        sec.top_margin = sec.bottom_margin = Inches(0.7)
        self.text_width = sec.page_width - sec.left_margin - sec.right_margin
        normal = self.doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10)
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.space_before = Pt(0)

    def para(self, space_before=0, space_after=0):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        return p

    # -- blocks ---------------------------------------------------------
    def header(self, node):
        h1 = node.find("h1")
        p = self.para(space_after=2)
        run = p.add_run(squash(h1.text()))
        run.bold, run.font.size, run.font.color.rgb = True, Pt(20), INK
        run.font.name = "Cambria"
        last = p
        for el in node.kids("p"):
            last = self.para()
            inline(last, el, {"color": MUTED, "size": 9.5})
        last.paragraph_format.space_after = Pt(4)
        bottom_border(last, size=12)

    def h2(self, node):
        p = self.para(space_before=12, space_after=4)
        run = p.add_run(squash(node.text()).upper())
        run.bold, run.font.size, run.font.color.rgb = True, Pt(10.5), ACCENT
        bottom_border(p, color="DFE4EC", size=6)
        p.paragraph_format.keep_with_next = True

    def sub(self, node):
        p = self.para(space_before=6, space_after=2)
        run = p.add_run(squash(node.text()).upper())
        run.bold, run.font.size, run.font.color.rgb = True, Pt(9), MUTED
        p.paragraph_format.keep_with_next = True

    def group(self, node):
        p = self.para(space_before=10, space_after=2)
        run = p.add_run(squash(node.text()))
        run.bold, run.font.size, run.font.color.rgb = True, Pt(10), INK
        p.paragraph_format.keep_with_next = True

    def title_row(self, h3, when, space_before=4):
        p = self.para(space_before=space_before)
        p.paragraph_format.tab_stops.add_tab_stop(self.text_width, WD_TAB_ALIGNMENT.RIGHT)
        p.paragraph_format.keep_with_next = True
        if h3 is not None:
            inline(p, h3, {"b": True, "color": INK})
        if when is not None:
            run = p.add_run("\t" + squash(when.text()))
            run.font.color.rgb = MUTED
            run.font.size = Pt(9.5)
        return p

    def item(self, node):
        when = node.find("span", "when")
        h3 = node.find("h3")
        self.title_row(h3, when)
        holder = node.kids("div")[0] if "row" in node.classes and node.kids("div") else node
        for el in holder.kids():
            if el.tag == "p":
                p = self.para()
                inline(p, el, {"color": MUTED})
            elif el.tag == "ul":
                for li in el.kids("li"):
                    p = self.doc.add_paragraph(style="List Bullet")
                    p.paragraph_format.space_after = Pt(2)
                    inline(p, li, {"color": MUTED})

    def skills(self, node):
        dts, dds = node.kids("dt"), node.kids("dd")
        for dt, dd in zip(dts, dds):
            p = self.para(space_after=2)
            run = p.add_run(squash(dt.text()) + ": ")
            run.bold, run.font.color.rgb = True, INK
            inline(p, dd, {"color": MUTED})

    def pubs(self, node):
        items = node.kids("li")
        n = len(items)
        for i, li in enumerate(items):
            p = self.para(space_after=4)
            fmt = p.paragraph_format
            fmt.left_indent, fmt.first_line_indent = Inches(0.4), Inches(-0.4)
            fmt.tab_stops.add_tab_stop(Inches(0.4))
            run = p.add_run(f"{n - i}.\t")
            run.bold, run.font.color.rgb, run.font.size = True, ACCENT, Pt(9)
            inline(p, li, {"color": MUTED, "size": 9})

    def block(self, node):
        cls = node.classes
        if node.tag == "header":
            self.header(node)
        elif node.tag == "section" or (node.tag == "div" and "cv-talks" in cls):
            for c in node.kids():
                self.block(c)
        elif node.tag == "h2":
            self.h2(node)
        elif node.tag == "h3" and "group" in cls:
            self.group(node)
        elif node.tag == "h3" and "sub" in cls:
            self.sub(node)
        elif node.tag == "div" and "item" in cls:
            self.item(node)
        elif node.tag == "dl":
            self.skills(node)
        elif node.tag == "ol" and "cv-pubs" in cls:
            self.pubs(node)
        elif node.tag == "p":
            p = self.para(space_before=2 if "updated" not in cls else 14, space_after=4)
            inline(p, node, {"color": MUTED if ("note" in cls or "updated" in cls or "keywords" in cls) else INK,
                             "size": 8.5 if "updated" in cls else None})


def html_to_docx(html_text, out_path, title="Curriculum Vitae"):
    tb = TreeBuilder()
    tb.feed(html_text)
    sheet = tb.root.find("main", "sheet")
    if sheet is None:
        raise ValueError("cv.html has no <main class='sheet'>")
    w = Writer()
    w.doc.core_properties.title = title
    for node in sheet.kids():
        w.block(node)
    w.doc.save(str(out_path))
